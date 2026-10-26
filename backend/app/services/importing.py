"""Safe import of contacts, send history and suppression lists (XLSX / CSV).

Safety properties
* The upload is treated as hostile: size cap, zip-bomb check, no formula evaluation, formula-looking
  cells are stored as inert text (exports neutralise them), control/invisible characters are stripped.
* Idempotent + resumable: rows are processed in chunks; `import_batches.rows_processed` is the
  checkpoint; unique constraints make a repeated run insert nothing twice.
* A contacts import can NEVER lift a suppression, re-enable an excluded contact, or overwrite an
  existing contact's eligibility/history. Existing contacts are left exactly as they are.
* Imported contacts start with history = 'unknown' unless the operator explicitly attests otherwise.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from collections import Counter
from collections.abc import Callable, Iterator
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from app.config import get_settings
from app.core.events import audit, record_event
from app.core.text import clean_text, email_domain, normalize_email, tld_country
from app.db import session_scope
from app.models import (
    Company,
    Contact,
    EligibilityDecision,
    ImportBatch,
    ImportRowIssue,
    OutreachEvent,
)
from app.services import jurisdictions, suppression
from app.services.eligibility import classify_new_contact

MAX_UPLOAD_BYTES = 60 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 500 * 1024 * 1024
MAX_ROWS = 500_000
CHUNK = 1000
ALLOWED_EXT = {".xlsx", ".csv"}

FIELDS: dict[str, list[tuple[str, bool, str]]] = {
    # (field, required, description)
    "contacts": [
        ("email", True, "Email address"),
        ("name", False, "Person's name"),
        ("company", False, "Company name"),
        ("source_row", False, "Row number in your original file"),
        ("country", False, "Country (2-letter code)"),
        ("title", False, "Job title"),
    ],
    "history": [
        ("email", True, "Email address that was contacted"),
        ("sent_at", False, "When it was sent"),
        ("status", False, "Outcome: sent / bounced / replied / unsubscribed / unknown"),
        ("note", False, "Free-text note"),
    ],
    "suppressions": [
        ("value", True, "Email address or domain to never contact"),
        ("reason", False, "Why (free text)"),
    ],
}

_SYNONYMS = {
    "email": ["email", "e-mail", "mail", "email address", "emailaddress", "work email"],
    "name": ["name", "full name", "fullname", "contact", "contact name", "person"],
    "company": ["company", "company name", "organisation", "organization", "employer", "account"],
    "source_row": ["source_row", "row", "source row", "original row"],
    "country": ["country", "country code", "iso country"],
    "title": ["title", "job title", "position", "role"],
    "sent_at": ["sent_at", "sent", "date", "sent date", "sent on", "timestamp", "date sent"],
    "status": ["status", "outcome", "result", "state"],
    "note": ["note", "notes", "comment", "comments"],
    "value": ["value", "email", "domain", "address"],
    "reason": ["reason", "why", "note"],
}


class ImportError_(ValueError):
    """User-facing import problem (bad file, bad mapping)."""


# --- file handling --------------------------------------------------------------------------------------------


def safe_filename(name: str) -> str:
    base = Path(name).name
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", base).strip("._") or "upload"
    return base[:120]


def save_upload(content: bytes, filename: str) -> tuple[Path, str]:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise ImportError_(f"Unsupported file type '{ext}'. Upload .xlsx or .csv")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ImportError_(f"File is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
    if not content:
        raise ImportError_("The file is empty")
    if ext == ".xlsx":
        _check_xlsx_zip(content)
    sha = hashlib.sha256(content).hexdigest()
    uploads = get_settings().uploads_dir
    uploads.mkdir(parents=True, exist_ok=True)
    path = uploads / f"{sha[:16]}-{safe_filename(filename)}"
    if not path.exists():
        path.write_bytes(content)
    return path, sha


def _check_xlsx_zip(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            total = sum(i.file_size for i in z.infolist())
            if total > MAX_UNCOMPRESSED_BYTES:
                raise ImportError_("Spreadsheet expands to an unreasonable size (possible zip bomb)")
            names = {i.filename for i in z.infolist()}
            if "xl/workbook.xml" not in names:
                raise ImportError_("Not a valid .xlsx workbook")
    except zipfile.BadZipFile as e:
        raise ImportError_("Not a valid .xlsx file") from e


def _decode_csv(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ImportError_("Could not decode the CSV file")


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def sheet_names(path: Path) -> list[str]:
    if path.suffix.lower() != ".xlsx":
        return []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as e:  # noqa: BLE001
        raise ImportError_(f"The file could not be read safely ({type(e).__name__})") from e
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def iter_raw_rows(path: Path, sheet: str | None = None) -> Iterator[tuple[int, list[Any]]]:
    """Yield (physical_row_number, values). Row 1 is the header."""
    if path.suffix.lower() == ".xlsx":
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                yield i, list(row)
        finally:
            wb.close()
    else:
        text = _decode_csv(path)
        delim = _sniff_delimiter(text)
        csv.field_size_limit(10_000_000)
        for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delim), start=1):
            yield i, row


def _guarded(it: Iterator[tuple[int, list[Any]]]) -> Iterator[tuple[int, list[Any]]]:
    """Any parser failure on a hostile or corrupt file becomes a clear, user-facing error (never a 500)."""
    try:
        yield from it
    except ImportError_:
        raise
    except Exception as e:  # noqa: BLE001 - zip/XML/encoding errors, defusedxml refusals, corrupt workbooks
        raise ImportError_(f"The file could not be read safely ({type(e).__name__}). Re-save it as a plain .xlsx or .csv and try again") from e


def read_table(path: Path, sheet: str | None = None) -> tuple[list[str], Iterator[tuple[int, dict[str, Any]]]]:
    it = _guarded(iter_raw_rows(path, sheet))
    try:
        _, header_row = next(it)
    except StopIteration:
        raise ImportError_("The file has no rows")
    headers: list[str] = []
    seen: Counter[str] = Counter()
    for j, h in enumerate(header_row):
        name = clean_text(h, 120) or f"column_{j + 1}"
        seen[name] += 1
        headers.append(name if seen[name] == 1 else f"{name}_{seen[name]}")

    def rows() -> Iterator[tuple[int, dict[str, Any]]]:
        for n, values in it:
            if not values or all(v is None or str(v).strip() == "" for v in values):
                continue
            yield n, {headers[k]: (values[k] if k < len(values) else None) for k in range(len(headers))}

    return headers, rows()


def suggest_mapping(kind: str, headers: list[str]) -> dict[str, str | None]:
    lowered = {h.lower().strip(): h for h in headers}
    mapping: dict[str, str | None] = {}
    used: set[str] = set()
    for field, _, _ in FIELDS[kind]:
        found = None
        for syn in _SYNONYMS.get(field, [field]):
            h = lowered.get(syn)
            if h and h not in used:
                found = h
                break
        mapping[field] = found
        if found:
            used.add(found)
    return mapping


def preview(path: Path, kind: str, sheet: str | None = None, sample: int = 12) -> dict[str, Any]:
    headers, rows = read_table(path, sheet)
    sample_rows: list[dict[str, Any]] = []
    total = 0
    for _, row in rows:
        total += 1
        if len(sample_rows) < sample:
            sample_rows.append({k: _preview_value(v) for k, v in row.items()})
        if total > MAX_ROWS:
            raise ImportError_(f"More than {MAX_ROWS:,} rows; split the file")
    mapping = suggest_mapping(kind, headers)
    mapped = {v for v in mapping.values() if v}
    extras = [h for h in headers if h not in mapped and h.lower() in {"data_quality_issue"}]
    return {
        "sheets": sheet_names(path),
        "headers": headers,
        "sample_rows": sample_rows,
        "total_rows": total,
        "suggested_mapping": mapping,
        "suggested_extra_columns": extras,
        "fields": [{"field": f, "required": r, "description": d} for f, r, d in FIELDS[kind]],
    }


def _preview_value(v: Any) -> Any:
    if v is None or isinstance(v, (int, float, bool)):
        return v
    if isinstance(v, datetime):
        return v.isoformat()
    return str(v)[:200]


def validate_mapping(kind: str, headers: list[str], mapping: dict[str, str | None]) -> None:
    valid = set(headers)
    for field, required, _ in FIELDS[kind]:
        col = mapping.get(field)
        if required and not col:
            raise ImportError_(f"Map a column to '{field}' (required)")
        if col and col not in valid:
            raise ImportError_(f"Column '{col}' mapped to '{field}' does not exist in the file")
    chosen = [c for c in mapping.values() if c]
    if len(chosen) != len(set(chosen)):
        raise ImportError_("The same column is mapped to more than one field")


# --- processing ------------------------------------------------------------------------------------------------------


def _get(row: dict[str, Any], mapping: dict[str, str | None], field: str) -> Any:
    col = mapping.get(field)
    return row.get(col) if col else None


def _parse_dt(v: Any) -> datetime | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v if v.tzinfo else v.replace(tzinfo=timezone.utc)
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d/%m/%Y %H:%M"):
        try:
            return datetime.strptime(s[:19] if "T" in s else s, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return d if d.tzinfo else d.replace(tzinfo=timezone.utc)
    except ValueError:
        return None


_HISTORY_STATUS = {
    "sent": "sent", "delivered": "sent", "submitted": "sent", "ok": "sent", "success": "sent", "contacted": "sent",
    "bounced": "bounced", "bounce": "bounced", "hard bounce": "bounced", "failed": "bounced", "invalid": "bounced",
    "replied": "replied", "reply": "replied", "responded": "replied", "answered": "replied",
    "unsubscribed": "unsubscribed", "unsubscribe": "unsubscribed", "opt-out": "unsubscribed", "optout": "unsubscribed",
    "opted out": "unsubscribed", "dnc": "unsubscribed", "do not contact": "unsubscribed", "complaint": "unsubscribed",
    "spam": "unsubscribed",
}


def _history_status(raw: Any) -> str:
    s = (clean_text(raw, 60) or "").lower()
    return _HISTORY_STATUS.get(s, "unknown")


class _Counts:
    def __init__(self) -> None:
        self.inserted = 0
        self.updated = 0
        self.duplicate_in_file = 0
        self.duplicate_existing = 0
        self.invalid = 0
        self.skipped = 0
        self.issues: list[ImportRowIssue] = []


def _upsert_companies(db: Session, domains_names: dict[str, str | None]) -> dict[str, int]:
    if not domains_names:
        return {}
    rows = [
        {"domain": d, "name": n, "country_code": tld_country(d)}
        for d, n in domains_names.items()
    ]
    stmt = pg_insert(Company).values(rows).on_conflict_do_nothing(index_elements=["domain"])
    db.execute(stmt)
    found = db.execute(select(Company.domain, Company.id).where(Company.domain.in_(list(domains_names)))).all()
    return {d: i for d, i in found}


def _build_contact_rows(
    db: Session,
    batch: ImportBatch,
    chunk: list[tuple[int, dict[str, Any]]],
    counts: _Counts,
    seen_in_file: set[str],
    history_default: str,
    force_history: str | None = None,
) -> tuple[list[dict], dict[str, str | None]]:
    mapping = batch.mapping
    extras: list[str] = batch.options.get("extra_columns", [])
    pending: dict[str, dict] = {}
    domains_names: dict[str, str | None] = {}
    for n, row in chunk:
        raw_email = _get(row, mapping, "email")
        email = normalize_email(raw_email)
        if email is None:
            counts.invalid += 1
            counts.issues.append(ImportRowIssue(batch_id=batch.id, source_row=n, issue="invalid_email",
                                                raw={"email": str(raw_email)[:200] if raw_email is not None else None}))
            continue
        if email in seen_in_file or email in pending:
            counts.duplicate_in_file += 1
            counts.issues.append(ImportRowIssue(batch_id=batch.id, source_row=n, issue="duplicate_in_file", raw={"email": email}))
            continue
        domain = email_domain(email)
        company_raw = clean_text(_get(row, mapping, "company"), 300)
        name = clean_text(_get(row, mapping, "name"), 300)
        src_row = _get(row, mapping, "source_row")
        try:
            src_row = int(src_row) if src_row not in (None, "") else n
        except (TypeError, ValueError):
            src_row = n
        status, reason, freemail, role = classify_new_contact(email, domain)
        country = clean_text(_get(row, mapping, "country"), 2)
        country_source = "import" if country else None
        if not country:
            country = tld_country(domain)
            country_source = "tld" if country else None
        else:
            country = country.upper()
        attrs = {}
        for col in extras:
            v = clean_text(row.get(col), 500)
            if v:
                attrs[col] = v
        pending[email] = dict(
            email_normalized=email, email_original=clean_text(raw_email, 320) or email, email_domain=domain, name=name,
            company_name_raw=company_raw, import_batch_id=batch.id, source_row=src_row, attributes=attrs,
            country_code=country, country_source=country_source,
            role_title=clean_text(_get(row, mapping, "title"), 200),
            eligibility_status=status, eligibility_reason=reason, is_freemail=freemail, is_role_account=role,
            history_status=force_history or history_default,
        )
        if not freemail:
            domains_names.setdefault(domain, company_raw)
            if company_raw and domains_names[domain] is None:
                domains_names[domain] = company_raw
    return list(pending.values()), domains_names


def _insert_contacts(db: Session, batch: ImportBatch, records: list[dict], domains_names: dict[str, str | None], counts: _Counts) -> list[tuple[int, str]]:
    if not records:
        return []
    emails = [r["email_normalized"] for r in records]
    existing = set(db.execute(select(Contact.email_normalized).where(Contact.email_normalized.in_(emails))).scalars())
    fresh = [r for r in records if r["email_normalized"] not in existing]
    counts.duplicate_existing += len(records) - len(fresh)
    for r in records:
        if r["email_normalized"] in existing:
            counts.issues.append(ImportRowIssue(batch_id=batch.id, source_row=r["source_row"], issue="already_in_database",
                                                raw={"email": r["email_normalized"]}))
    if not fresh:
        return []
    company_ids = _upsert_companies(db, domains_names)
    jurisdictions.ensure_policy_rows(db, {r["country_code"] for r in fresh if r["country_code"]})
    for r in fresh:
        r["company_id"] = company_ids.get(r["email_domain"])
        r["eligibility_updated_at"] = datetime.now(timezone.utc)
    stmt = pg_insert(Contact).values(fresh).on_conflict_do_nothing(index_elements=["email_normalized"]).returning(Contact.id, Contact.email_normalized)
    inserted = db.execute(stmt).all()
    counts.inserted += len(inserted)
    counts.duplicate_existing += len(fresh) - len(inserted)  # lost a race: someone else inserted it
    by_email = {r["email_normalized"]: r for r in fresh}
    decisions = []
    events = []
    for cid, email in inserted:
        r = by_email[email]
        decisions.append(dict(contact_id=cid, decided_by="import", field="eligibility_status", from_value=None,
                              to_value=r["eligibility_status"], reason=r["eligibility_reason"] or f"imported (batch {batch.id})"))
        decisions.append(dict(contact_id=cid, decided_by="import", field="history_status", from_value=None,
                              to_value=r["history_status"], reason=f"import batch {batch.id}: {batch.history_attestation}"))
        events.append(dict(event_type="imported", contact_id=cid, source="system", actor=batch.created_by,
                           data={"batch_id": batch.id, "source_row": r["source_row"]}))
    if decisions:
        db.execute(pg_insert(EligibilityDecision).values(decisions))
    if events:
        db.execute(pg_insert(OutreachEvent).values(events))
    return [(cid, email) for cid, email in inserted]


def _process_contacts_chunk(db: Session, batch: ImportBatch, chunk, counts: _Counts, seen: set[str]) -> None:
    history_default = "attested_unsent" if batch.history_attestation == "attested_unsent" else "unknown"
    records, domains_names = _build_contact_rows(db, batch, chunk, counts, seen, history_default)
    _insert_contacts(db, batch, records, domains_names, counts)
    seen.update(r["email_normalized"] for r in records)


def _process_history_chunk(db: Session, batch: ImportBatch, chunk, counts: _Counts, seen: set[str]) -> None:
    mapping = batch.mapping
    # make sure every emailed address exists as a contact (history must never be lost)
    base_chunk = [(n, {"__email": _get(row, mapping, "email")}) for n, row in chunk]
    fake_batch_mapping = {"email": "__email"}
    tmp = ImportBatch(id=batch.id, mapping=fake_batch_mapping, options={}, history_attestation="unknown", created_by=batch.created_by)
    records, domains_names = _build_contact_rows(db, tmp, base_chunk, _Counts(), set(), "previously_contacted", "previously_contacted")
    _insert_contacts(db, batch, records, domains_names, _Counts())

    for n, row in chunk:
        email = normalize_email(_get(row, mapping, "email"))
        if email is None:
            counts.invalid += 1
            counts.issues.append(ImportRowIssue(batch_id=batch.id, source_row=n, issue="invalid_email", raw={}))
            continue
        contact = db.execute(select(Contact).where(Contact.email_normalized == email)).scalar_one()
        status = _history_status(_get(row, mapping, "status"))
        sent_at = _parse_dt(_get(row, mapping, "sent_at"))
        note = clean_text(_get(row, mapping, "note"), 500)
        before = contact.history_status
        if before != "previously_contacted":
            contact.history_status = "previously_contacted"
            db.add(EligibilityDecision(contact_id=contact.id, decided_by="history_import", field="history_status",
                                       from_value=before, to_value="previously_contacted",
                                       reason=f"sent-history import batch {batch.id} (row {n}, status {status})"))
            counts.updated += 1
        else:
            counts.duplicate_existing += 1
        ev = OutreachEvent(event_type="historical_send", contact_id=contact.id, source="system", actor=batch.created_by,
                           data={"batch_id": batch.id, "row": n, "status": status, "note": note,
                                 "sent_at": sent_at.isoformat() if sent_at else None})
        if sent_at:
            ev.occurred_at = sent_at
        db.add(ev)
        if status == "bounced":
            suppression.suppress(db, value=email, reason="bounce_hard", source=f"history import batch {batch.id}",
                                 contact_id=contact.id, actor=batch.created_by)
        elif status == "unsubscribed":
            suppression.suppress(db, value=email, reason="unsubscribe", source=f"history import batch {batch.id}",
                                 contact_id=contact.id, actor=batch.created_by)
        elif status == "replied" and contact.pipeline_stage in ("new", "researched", "qualified", "draft_ready", "approved", "contacted"):
            record_event(db, "stage_changed", contact_id=contact.id, source="system", actor=batch.created_by,
                         data={"from": contact.pipeline_stage, "to": "replied", "why": "history import"})
            contact.pipeline_stage = "replied"
        elif status in ("sent", "unknown") and contact.pipeline_stage in ("new", "researched", "qualified", "draft_ready", "approved"):
            contact.pipeline_stage = "contacted"
        db.flush()


def _process_suppressions_chunk(db: Session, batch: ImportBatch, chunk, counts: _Counts, seen: set[str]) -> None:
    mapping = batch.mapping
    for n, row in chunk:
        raw = clean_text(_get(row, mapping, "value"), 320)
        note = clean_text(_get(row, mapping, "reason"), 300)
        if not raw:
            counts.invalid += 1
            continue
        kind = "email" if "@" in raw else "domain"
        try:
            suppression.suppress(db, kind=kind, value=raw, reason="import", source=f"suppression import batch {batch.id}",
                                 actor=batch.created_by, notes=note)
            counts.inserted += 1
        except suppression.SuppressionError:
            counts.invalid += 1
            counts.issues.append(ImportRowIssue(batch_id=batch.id, source_row=n, issue="invalid_value", raw={"value": raw[:200]}))


def create_batch(
    db: Session,
    *,
    path: Path,
    sha256: str,
    filename: str,
    kind: str,
    mapping: dict[str, str | None],
    options: dict[str, Any] | None = None,
    sheet: str | None = None,
    provenance_note: str | None = None,
    history_attestation: str = "unknown",
    created_by: str | None = None,
) -> ImportBatch:
    """Validate the mapping against the real file, count rows, and register a batch (not yet running)."""
    if kind not in FIELDS:
        raise ImportError_(f"unknown import kind {kind!r}")
    if history_attestation not in ("unknown", "attested_unsent"):
        raise ImportError_("history_attestation must be 'unknown' or 'attested_unsent'")
    if kind == "contacts" and history_attestation == "attested_unsent" and not (provenance_note and len(provenance_note.strip()) >= 10):
        raise ImportError_("To attest that none of these contacts were emailed before, describe the list's origin and "
                           "how you know (10+ characters)")
    headers, rows = read_table(path, sheet)
    validate_mapping(kind, headers, mapping)
    total = 0
    for _ in rows:
        total += 1
        if total > MAX_ROWS:
            raise ImportError_(f"More than {MAX_ROWS:,} rows; split the file")
    batch = ImportBatch(
        kind=kind, filename=filename[:300], file_sha256=sha256, file_path=str(path), status="uploaded",
        sheet_name=sheet, mapping={k: v for k, v in mapping.items() if v}, options=options or {},
        provenance_note=(provenance_note or "").strip() or None, history_attestation=history_attestation,
        total_rows=total, created_by=created_by,
    )
    db.add(batch)
    db.flush()
    audit(db, "import.create", actor=created_by, entity_type="import_batch", entity_id=batch.id,
          details={"kind": kind, "filename": filename, "rows": total, "history_attestation": history_attestation,
                   "provenance_note": batch.provenance_note})
    return batch


_PROCESSORS: dict[str, Callable] = {
    "contacts": _process_contacts_chunk,
    "history": _process_history_chunk,
    "suppressions": _process_suppressions_chunk,
}


def process_batch(
    batch_id: int,
    progress: Callable[[int, int], None] | None = None,
    should_cancel: Callable[[], bool] | None = None,
) -> dict[str, Any]:
    """Run (or resume) an import. Safe to call again after a crash: it continues from rows_processed."""
    with session_scope() as db:
        batch = db.get(ImportBatch, batch_id)
        if batch is None:
            raise ImportError_("import batch not found")
        if batch.status == "completed":
            return {"status": "completed", "already": True}
        batch.status = "running"
        path = Path(batch.file_path)
        kind, mapping, sheet = batch.kind, batch.mapping, batch.sheet_name
        start_at = batch.rows_processed
        total = batch.total_rows
    headers, rows = read_table(path, sheet)
    validate_mapping(kind, headers, mapping)
    processor = _PROCESSORS[kind]

    processed = 0
    seen: set[str] = set()
    chunk: list[tuple[int, dict[str, Any]]] = []

    def flush(chunk_rows: list[tuple[int, dict[str, Any]]], upto: int) -> None:
        with session_scope() as db:
            b = db.get(ImportBatch, batch_id)
            counts = _Counts()
            processor(db, b, chunk_rows, counts, seen)
            b.inserted += counts.inserted
            b.updated += counts.updated
            b.duplicate_in_file += counts.duplicate_in_file
            b.duplicate_existing += counts.duplicate_existing
            b.invalid += counts.invalid
            b.rows_processed = upto
            db.add_all(counts.issues[:2000])
        if progress:
            progress(upto, total)

    cancelled = False
    for n, row in rows:
        processed += 1
        if processed <= start_at:
            # rows done in an earlier run: remember emails so in-file duplicates are still detected
            e = normalize_email(_get(row, mapping, "email" if kind != "suppressions" else "value"))
            if e:
                seen.add(e)
            continue
        chunk.append((n, row))
        if len(chunk) >= CHUNK:
            flush(chunk, processed)
            chunk = []
            if should_cancel and should_cancel():
                cancelled = True
                break
    if chunk and not cancelled:
        flush(chunk, processed)

    with session_scope() as db:
        b = db.get(ImportBatch, batch_id)
        if cancelled:
            b.status = "cancelled"
        else:
            b.status = "completed"
            b.finished_at = datetime.now(timezone.utc)
        audit(db, f"import.{'cancelled' if cancelled else 'completed'}", actor=b.created_by, entity_type="import_batch",
              entity_id=b.id, details={"kind": b.kind, "inserted": b.inserted, "updated": b.updated,
                                       "duplicate_in_file": b.duplicate_in_file, "duplicate_existing": b.duplicate_existing,
                                       "invalid": b.invalid})
        result = {
            "status": b.status, "kind": b.kind, "rows": b.rows_processed, "inserted": b.inserted, "updated": b.updated,
            "duplicate_in_file": b.duplicate_in_file, "duplicate_existing": b.duplicate_existing, "invalid": b.invalid,
        }
    return result
